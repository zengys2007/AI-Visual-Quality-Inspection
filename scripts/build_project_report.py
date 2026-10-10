"""生成《电池包瑕疵检测》竞赛项目方案报告 Word 文档。"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "电池包瑕疵检测_项目方案报告.docx"

NAVY = RGBColor(0x1A, 0x36, 0x5D)
ACCENT = RGBColor(0x16, 0x77, 0xFF)
INK = RGBColor(0x1A, 0x23, 0x32)
MUTED = RGBColor(0x5C, 0x6B, 0x7A)


def set_run_font(run, name="宋体", size=12, bold=False, color=INK, east="宋体"):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = name
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:eastAsia"), east)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.35
    sizes = {1: 16, 2: 14, 3: 12}
    run = p.add_run(text)
    set_run_font(run, "黑体", sizes.get(level, 12), True, NAVY if level == 1 else ACCENT, "黑体")
    return p


def add_body(doc, text, first_line=True):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(2)
    pf.space_after = Pt(6)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if first_line:
        pf.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    set_run_font(run, "宋体", 12, False, INK, "宋体")
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(text)
    set_run_font(run, "楷体", 10.5, False, MUTED, "楷体")
    return p


def shade_header_row(row, fill="1A365D"):
    for cell in row.cells:
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = tcPr.makeelement(
            qn("w:shd"),
            {
                qn("w:val"): "clear",
                qn("w:color"): "auto",
                qn("w:fill"): fill,
            },
        )
        tcPr.append(shd)


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        set_run_font(run, "黑体", 10.5, True, RGBColor(255, 255, 255), "黑体")
    shade_header_row(hdr)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(val))
            set_run_font(run, "宋体", 10.5, False, INK, "宋体")
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return table


def set_cell_text(cell, text, size=12, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, east="宋体"):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    run = p.add_run(text)
    set_run_font(run, "黑体" if bold else "宋体", size, bold, INK, east)


def build():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.8)
        section.right_margin = Cm(2.6)
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)

    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    # Cover
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("绍兴市第一届大学生人工智能应用竞赛")
    set_run_font(run, "黑体", 16, True, NAVY, "黑体")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("智能制造与工业智能 AI 应用赛道")
    set_run_font(run, "黑体", 14, True, ACCENT, "黑体")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(28)
    run = p.add_run("项目方案报告")
    set_run_font(run, "黑体", 26, True, NAVY, "黑体")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    run = p.add_run("赛题二：电池包瑕疵检测")
    set_run_font(run, "黑体", 18, True, ACCENT, "黑体")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("（表面裂纹 Surface_Crack 视觉质检系统）")
    set_run_font(run, "楷体", 14, False, MUTED, "楷体")

    meta = [
        ("作品名称", "电池包表面裂纹视觉质检系统（AI-Visual-Quality-Inspection）"),
        ("所选赛题", "赛题二：电池包瑕疵检测"),
        ("团队名称", "视觉检测团队（可按报名信息修改）"),
        ("仓库工程", "AI-Visual-Quality-Inspection"),
        ("文档版本", "v1.0（依据赛题说明书、附件1及仓库实现整理）"),
    ]
    doc.add_paragraph()
    table = doc.add_table(rows=len(meta), cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (k, v) in enumerate(meta):
        set_cell_text(table.rows[i].cells[0], k, 12, True)
        set_cell_text(table.rows[i].cells[1], v, 12, False, WD_ALIGN_PARAGRAPH.LEFT)
        table.rows[i].cells[0].width = Cm(3.6)
        table.rows[i].cells[1].width = Cm(12.4)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(36)
    run = p.add_run("说明：本报告内容整合自竞赛赛题文档、附件1输出规范、仓库 docs/ 接口规范与训练验证预测说明，以及 src/、app/、configs/ 实现。")
    set_run_font(run, "楷体", 10.5, False, MUTED, "楷体")

    doc.add_page_break()

    # TOC-like outline
    add_heading(doc, "目录", 1)
    toc = [
        "一、研究现状与意义",
        "二、目标与解决思路",
        "三、算法设计与实现",
        "四、算法结果与结果分析",
        "五、可视化应用与工程交付",
        "六、团队介绍及分工",
        "七、总结与后续工作",
    ]
    for t in toc:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.8
        p.paragraph_format.left_indent = Cm(1)
        run = p.add_run(t)
        set_run_font(run, "宋体", 14, False, INK, "宋体")

    doc.add_page_break()

    add_heading(doc, "一、研究现状与意义", 1)
    add_heading(doc, "1.1 竞赛与产业背景", 2)
    add_body(
        doc,
        "随着“中国制造2025”战略深入推进，智能制造已成为产业升级的核心方向。在电子、新能源、纺织等精密制造领域，产品质量直接关系到产品性能、安全性和品牌声誉。传统人工目检存在效率低、成本高、主观性强、易疲劳等固有缺陷，难以满足现代化产线对高通量、高精度、高稳定性的质检需求。",
    )
    add_body(
        doc,
        "机器视觉与深度学习的融合，为工业质检提供了可规模化的技术路径：通过对产品表面瑕疵进行毫秒级识别与定位，可提升检测效率与一致性，降低不良品流出风险。但工业场景中的瑕疵往往形态多样、对比度低、背景复杂，对算法鲁棒性、泛化能力与工程落地能力提出了很高要求。本赛题正是面向真实工业质检需求，要求参赛团队交付“可运行、可复现、可评估”的完整视觉检测工程。",
    )

    add_heading(doc, "1.2 电池包表面裂纹检测的业务意义", 2)
    add_body(
        doc,
        "新能源电池产业快速发展，电池包在装配、焊接、连接、老化与出厂检测等环节对外观一致性、安全性和可追溯性的要求持续提升。表面裂纹可能影响导电可靠性、绝缘安全、后续装配质量及整包寿命。若依赖人工目检，容易受到经验、疲劳、光照与节拍波动影响，难以在高通量产线中保持稳定一致的检测效果。",
    )
    add_body(
        doc,
        "因此，围绕电池包表面裂纹（赛题规定类别名称为 Surface_Crack）构建自动化视觉检测系统，具有直接的安全价值与产线落地价值：一方面降低漏检带来的质量与安全风险，另一方面将检测结果结构化输出，便于质量追溯与评分工具统一评价。",
    )

    add_heading(doc, "1.3 技术现状与本赛题挑战", 2)
    add_body(
        doc,
        "工业表面缺陷检测近年来主要沿两条技术路线发展：一是以 YOLO、Faster R-CNN 等为代表的有监督目标检测/实例分割；二是结合传统视觉（阈值、形态学、边缘与轮廓）或无监督异常检测以应对样本不均衡。电池包裂纹检测的共性难点包括：裂纹细长、对比度偏低、成像角度与光照变化大、正常样本与缺陷样本比例不均衡、小目标定位难。",
    )
    add_body(
        doc,
        "赛题说明书明确：算法内部可采用检测框或实例分割 Mask，但对外结果（界面展示框与 CSV）必须统一为轴对齐矩形框；分割结果须对每个独立裂纹实例计算外接矩形，坐标为原始图像像素绝对值。本项目据此选择“实例分割 + 外接矩形导出”的路线，既保留裂纹形态信息，又满足评分接口。",
    )

    add_heading(doc, "1.4 赛题要求摘要（仅电池包）", 2)
    add_body(
        doc,
        "本队选择赛题二，不涉及窗帘布类别（Hole、Ink、Broken_Filament）。核心任务是：基于主办方统一数据集，识别电池包表面裂纹，输出缺陷类别、位置区域与置信度，并完成可视化应用与批量 CSV 导出。",
    )
    add_table(
        doc,
        ["任务", "检测目标", "输出要求"],
        [["表面裂纹检测", "识别电池包表面裂纹瑕疵", "类别 Surface_Crack、位置框、置信度；无缺陷输出 Ok 占位行"]],
        [3.2, 4.8, 7.8],
    )
    add_caption(doc, "表1 赛题二任务要求（依据赛题说明书表3-2-1整理）")

    add_table(
        doc,
        ["指标", "含义", "计算要点"],
        [
            ["表面裂纹识别率", "正确识别裂纹状态的图像占比", "正确识别数量 / 测试图像总数"],
            ["误检率", "OK 被误判为 NG 的比例", "误判 NG 的 OK 数 / 实际 OK 总数"],
            ["一致性准确率", "类别与位置同时正确的图像占比", "综合正确图像数 / 测试图像总数"],
            ["mAP50", "定位能力（IoU=0.5）", "预测框与真实框交并比评价"],
            ["Precision / Recall", "精确率 / 召回率", "正确检出 / 全部检出；正确检出 / 全部真值"],
        ],
        [3.6, 5.0, 7.2],
    )
    add_caption(doc, "表2 量化评价指标（依据赛题说明书表3-2-2整理）")

    add_heading(doc, "二、目标与解决思路", 1)
    add_heading(doc, "2.1 总体目标", 2)
    add_body(doc, "围绕竞赛评分中的算法性能、应用开发与创新应用三部分，本项目设定如下目标：")
    add_body(
        doc,
        "（1）算法：构建可训练、可验证、可推理的表面裂纹实例分割模型，在验证集上报告 Box/Mask 的 mAP、并选定合适置信度阈值平衡精确率与召回率。",
    )
    add_body(
        doc,
        "（2）接口：检测结果严格符合《附件1：检测结果输出表格规范及字段说明》与仓库 docs/接口规范.md：八列表头固定、类别仅 Surface_Crack / Ok、无缺陷必须一行占位、坐标为像素绝对值。",
    )
    add_body(
        doc,
        "（3）工程：交付可运行的桌面可视化应用，支持导入、检测、可视化、选中导出与选中清空；与算法脚本共同构成完整参赛作品。",
    )
    add_body(
        doc,
        "（4）规范：不得获取或反推测试集标签，不得人工改写测试结果规避统一评价。",
    )

    add_heading(doc, "2.2 解决思路", 2)
    add_body(
        doc,
        "总体思路是“数据规范化 → 实例分割训练 → Mask 转外接框 → 规范 CSV → 桌面应用交付”。具体如下。",
    )
    add_body(
        doc,
        "第一，数据侧：使用 Label Studio 按 YOLO 多边形/分割格式标注裂纹，经 src/organize_yolo_data.py 去除 UUID 前缀并复制到 Ultralytics 标准目录 data/yolo/images|labels/{train,val}。类别在 configs/data.yaml 中仅定义 0: Surface_Crack。当前仓库中训练集约 19 对图标、验证集 160 对图标；data/train 保留约 640 张原始训练图库供增量标注拷贝。",
    )
    add_body(
        doc,
        "第二，模型侧：采用 Ultralytics YOLO 实例分割（预训练权重 yolo26n-seg.pt），输入尺寸 640，训练 100 epoch，输出 best.pt。推理时以实例 Mask/多边形描述裂纹，再取轴对齐外接矩形写入 CSV，兼顾分割精度与评分接口。",
    )
    add_body(
        doc,
        "第三，阈值侧：预测脚本默认置信度阈值 CONF=0.223，取自验证曲线 F1 最优点附近，用于抑制低置信误检、同时尽量保留细裂纹召回。",
    )
    add_body(
        doc,
        "第四，应用侧：以 pywebview 套壳本地前端（app/desktop.py + app/static），实现拖放导入、列表多选/全选、单项删除、选中后导出 CSV、选中后清空。导出字段与 predict.py 一致，保证“界面展示与提交评分同一套记录规则”。",
    )

    add_heading(doc, "2.3 技术路线图", 2)
    add_table(
        doc,
        ["阶段", "输入", "关键模块", "输出"],
        [
            ["数据整理", "Label Studio YOLO 多边形 txt + data/train 原图", "organize_yolo_data.py", "data/yolo 标准划分"],
            ["训练", "configs/data.yaml + yolo26n-seg.pt", "src/train.py", "runs/.../best.pt"],
            ["验证", "val 图与多边形标签", "src/val.py", "Box/Mask mAP 与曲线图"],
            ["推理导出", "测试/验证图像 + best.pt", "src/predict.py", "八列 CSV + 可视化图 + txt"],
            ["桌面应用", "本地图像/文件夹", "app/desktop.py", "导入-检测-可视化-导出"],
        ],
        [2.6, 4.6, 3.6, 5.0],
    )
    add_caption(doc, "表3 工程技术路线（对应仓库 src/、app/、configs/）")

    add_heading(doc, "三、算法设计与实现", 1)
    add_heading(doc, "3.1 任务建模", 2)
    add_body(
        doc,
        "将电池包表面裂纹检测建模为单类别实例分割问题：输入一张工业图像，输出若干裂纹实例的分割轮廓及置信度。对外接口再将每个实例转换为轴对齐矩形 (x_min, y_min, x_max, y_max)。无任何实例时输出一行 Ok 占位，以保证测试图片总数可被评分工具统计。",
    )
    add_body(
        doc,
        "选择实例分割而非纯检测框的原因：裂纹形态细长、边界不规则，Mask 更能描述真实缺陷区域；赛题允许 Mask，但要求最终以最小外接矩形提交，因此分割是“对内更准、对外合规”的折中。",
    )

    add_heading(doc, "3.2 数据与标注格式", 2)
    add_body(
        doc,
        "数据集由主办方统一提供，覆盖不同外观、角度、光照、裂纹形态与程度，包含正常样本与裂纹样本。本仓库按 YOLO segment 官方结构组织：",
    )
    add_body(
        doc,
        "标注一行一个实例，格式为：class_id x1 y1 x2 y2 … xn yn，坐标归一化到 [0,1]。class_id=0 对应 Surface_Crack。空 txt 表示该图无裂纹目标。整理脚本识别 {uuid}-image_N.txt，匹配 data/train/image_N.jpg，写入 images/train 与 labels/train。验证集独立维护于 images/val、labels/val。",
        first_line=True,
    )
    add_table(
        doc,
        ["配置项", "取值"],
        [
            ["configs/data.yaml path", "data/yolo"],
            ["train / val", "images/train ， images/val"],
            ["names.0", "Surface_Crack"],
            ["当前 train 样本（仓库统计）", "19 张图 / 19 份标签"],
            ["当前 val 样本（仓库统计）", "160 张图 / 160 份标签"],
            ["原始图库 data/train", "约 640 张（整理脚本拷图来源）"],
        ],
        [7.0, 8.8],
    )
    add_caption(doc, "表4 数据集配置与规模（以仓库当前文件为准，可随增量标注更新）")

    add_heading(doc, "3.3 模型与训练实现", 2)
    add_body(
        doc,
        "训练入口为 src/train.py，封装 Ultralytics YOLO 官方训练流程。关键超参数：预训练模型 yolo26n-seg.pt（nano 分割模型，便于在有限算力下完成训练与演示）；epochs=100；imgsz=640；工程输出目录 runs/segment/surface_crack。Windows 下训练启动置于 if __name__ == \"__main__\" 内，避免多进程 RuntimeError；运行前切换到项目根目录，以保证 data.yaml 相对路径正确。有 GPU 时自动使用，否则 CPU。",
    )
    add_body(
        doc,
        "依赖见 requirements.txt：ultralytics==8.4.173，桌面端 pywebview>=5.0。可选先安装 CUDA 版 PyTorch 再安装本文件以启用 GPU。",
    )
    add_body(
        doc,
        "增量数据流程：新标注导出至 data/yolo/labels 根目录 → 运行 organize_yolo_data.py → 再运行 train.py → 将新的 best.pt 复制到项目根目录供 predict.py 使用。",
    )

    add_heading(doc, "3.4 验证实现", 2)
    add_body(
        doc,
        "src/val.py 读取 runs/segment/surface_crack/weights/best.pt（可改为根目录 best.pt），在 val 划分上评估 Box 与 Mask 指标，并输出 PR/F1 曲线、混淆矩阵与 batch 可视化。指标含义：mAP50 为 IoU=0.5 的平均精度，mAP75 更严格，mAP50-95 为 IoU 从 0.5 到 0.95 的平均。",
    )

    add_heading(doc, "3.5 推理、后处理与 CSV 导出", 2)
    add_body(
        doc,
        "src/predict.py 完成评分提交所需的批量推理。默认权重 best.pt，默认输入 data/yolo/images/val，置信度 0.223，imgsz=640，stream=True 以节省大量图片时的内存。每张图计算 preprocess+inference+postprocess 总和作为 inference_time_ms，同一张图多行填相同耗时。",
    )
    add_body(
        doc,
        "后处理规则与接口规范一致：若 boxes 为空，输出 [image_id, Ok, -1, -1, -1, -1, -1, t_ms]；若有框，则逐实例写出 Surface_Crack、保留两位小数的 confidence、以及 xyxy 像素坐标（脚本中 round 到 1 位小数）。同时写出与测试集同格式的多边形 txt（归一化，仅作内部/可视化，不得替代 CSV 四项框坐标）。",
    )
    add_table(
        doc,
        ["字段", "类型", "电池包填写规则"],
        [
            ["image_id", "字符串", "文件名含扩展名，与测试集逐字一致"],
            ["class_name", "字符串", "仅 Surface_Crack 或 Ok"],
            ["confidence", "浮点", "[0,1] 两位小数；Ok 为 -1"],
            ["x_min,y_min", "数", "左上角像素；Ok 为 -1"],
            ["x_max,y_max", "数", "右下角坐标而非宽高；须大于 min；Ok 为 -1"],
            ["inference_time_ms", "浮点", "该图整张检测总耗时，同图多行相同"],
        ],
        [4.2, 2.4, 9.2],
    )
    add_caption(doc, "表5 附件1 / 接口规范八列字段（窗帘布类别已忽略）")

    add_body(
        doc,
        "桌面端 app/static/app.js 以同一套规则生成 CSV：有裂纹只输出 Surface_Crack 行；无裂纹恰好一行 Ok；禁止窗帘布三类名称。desktop.py 通过 js_api.save_csv 弹出另存为对话框，UTF-8 写入“团队名称_检测结果.csv”。",
    )

    add_heading(doc, "3.6 可视化应用设计与实现", 2)
    add_body(
        doc,
        "赛题规定：未提供可运行可视化应用视为未完成本任务。本项目桌面壳基于 pywebview，加载 app/static/index.html。主要能力包括：",
    )
    add_body(doc, "（1）导入：按钮选择单张/文件夹，或从资源管理器、桌面拖入图像及文件夹（递归收集 jpg/jpeg/png/bmp/webp）。")
    add_body(doc, "（2）列表交互：勾选多选、全选/取消全选、条目右上角 × 删除单张、预览与勾选分离。")
    add_body(doc, "（3）检测与展示：单张/批量检测；预览叠加裂纹框与置信度；结果表字段对齐接口规范。当前桌面检测为前端 Mock，用于界面与导出流程演示；正式评分推理以 src/predict.py 的 YOLO 权重为准，二者 CSV 列定义保持一致。")
    add_body(doc, "（4）导出前置条件：必须先勾选图像，且选中项须已有检测结果，方可导出。")
    add_body(doc, "（5）清空前置条件：必须先勾选，仅删除选中项；全部清空需先全选再清空。")

    add_heading(doc, "四、算法结果与结果分析", 1)
    add_heading(doc, "4.1 验证集指标（仓库文档记录）", 2)
    add_body(
        doc,
        "docs/val.md 记录了一次实测参考：验证集 160 张图，权重 surface_crack/best.pt。结果如下。",
    )
    add_table(
        doc,
        ["指标", "Box（外接框）", "Mask（实例分割）"],
        [
            ["mAP50", "0.523", "0.409"],
            ["mAP50-95", "0.253", "0.106"],
            ["mAP75", "0.227", "0.010"],
        ],
        [5.2, 5.4, 5.2],
    )
    add_caption(doc, "表6 验证集实测参考（摘自 docs/val.md）")

    add_heading(doc, "4.2 结果分析", 2)
    add_body(
        doc,
        "（1）Box mAP50=0.523 高于 Mask mAP50=0.409，符合细长裂纹的几何特性：外接矩形对定位容错更大，而 Mask 要求像素级重合，在裂纹宽度仅数像素、标注多边形不完全贴合时，Mask 的 IoU 更容易下降。竞赛对外提交的是矩形框，Box mAP 与评分 mAP50 更直接相关。",
    )
    add_body(
        doc,
        "（2）mAP75 与 mAP50-95 明显低于 mAP50，说明高 IoU 下定位仍偏松。裂纹端点、分叉与涂覆纹理干扰会导致框偏大或偏移。后续可通过增加多边形标注精度、针对性增强（对比度、光照、随机旋转）、以及适当加大模型或微调 imgsz 改善。",
    )
    add_body(
        doc,
        "（3）Mask mAP75 仅 0.010，表明当前分割轮廓与真值在严格 IoU 下几乎不对齐。这对 CSV 外接框影响小于对 Mask 指标本身：只要外接矩形能包住裂纹主体，Box IoU=0.5 仍可能命中。工程上应继续用 Mask 辅助可视化，但优化目标需同时盯住 Box Precision/Recall。",
    )
    add_body(
        doc,
        "（4）训练集目前仅 19 张已整理样本，而验证集 160 张、原始图库约 640 张，数据瓶颈显著。小样本下 nano 分割模型容易欠拟合细裂纹、或过拟合已标注形态。这是当前指标的主要限制因素。优先路径是继续 Label Studio 标注并跑通增量整理—再训练闭环。",
    )
    add_body(
        doc,
        "（5）CONF=0.223 选在 F1 最优点附近，有利于在精确率与召回率之间折中：阈值过高会漏检浅裂纹（拉低 Recall 与裂纹识别率），过低会把涂覆纹理判为裂纹（拉高误检率）。提交测试集前应以 val 曲线再确认一次阈值。",
    )

    add_heading(doc, "4.3 与评分标准的对应", 2)
    add_table(
        doc,
        ["评分模块", "本科组权重", "专科组权重", "本方案对应交付"],
        [
            ["算法性能（mAP50/P/R）", "45%", "30%", "YOLO-seg 训练验证、阈值选择、规范框输出"],
            ["应用开发", "30%", "45%", "pywebview 桌面应用：导入/检测/可视化/选中导出"],
            ["创新应用", "25%", "25%", "分割转框、拖放导入、选中导出/清空、接口规范化"],
        ],
        [4.2, 3.0, 3.0, 5.6],
    )
    add_caption(doc, "表7 方案与赛题评分模块的对应关系")

    add_heading(doc, "4.4 误差与风险", 2)
    add_body(doc, "常见错误已在接口规范中列出，本项目在实现中针对性规避：")
    add_table(
        doc,
        ["风险", "后果", "本项目做法"],
        [
            ["坐标归一化", "IoU 失真", "CSV 使用像素绝对值；归一化多边形仅写内部 txt"],
            ["x_max/y_max 误填宽高", "坐标解析失败", "直接使用 xyxy 右下角"],
            ["同图耗时不一致", "平均推理时间错误", "按图汇总 speed 后各行复用"],
            ["无缺陷图缺行", "测试图总数遗漏", "强制输出一行 Ok 占位"],
            ["类别名不规范", "预测被忽略", "仅 Surface_Crack / Ok"],
            ["桌面 Mock 与真模型不一致", "演示与提交结果不同", "列定义对齐；正式提交走 predict.py"],
        ],
        [3.6, 3.6, 8.6],
    )
    add_caption(doc, "表8 风险控制（依据接口规范第4节与仓库实现）")

    add_heading(doc, "五、可视化应用与工程交付", 1)
    add_body(
        doc,
        "工程目录以仓库为准：算法脚本在 src/，说明在 docs/，数据配置在 configs/data.yaml，桌面应用在 app/。启动训练/验证/预测均在项目根目录调用对应脚本；桌面应用执行 python app/desktop.py（虚拟环境需已安装 pywebview 与 ultralytics）。",
    )
    add_body(
        doc,
        "提交评分 CSV 文件名格式为“团队名称_检测结果.csv”，示例为“视觉检测团队_检测结果.csv”，UTF-8 逗号分隔。预测产物目录示例：runs/segment/surface_crack_predict/ 下含 CSV、images 叠加图、txt 多边形。",
    )

    add_heading(doc, "六、团队介绍及分工", 1)
    add_heading(doc, "6.1 团队概况", 2)
    add_body(
        doc,
        "本队为 5 人小组，队名暂以仓库默认“视觉检测团队”填写，参赛时请替换为报名系统中的正式队名与队员姓名、学号、院校。小组按“数据—模型—工程—应用—文档测试”闭环分工，保证算法结果与可视化导出一致、符合附件1。",
    )

    add_heading(doc, "6.2 人员分工", 2)
    add_table(
        doc,
        ["角色", "姓名（请填写）", "主要职责", "对应仓库工作"],
        [
            [
                "队长 / 算法",
                "队员A",
                "总体技术路线、YOLO-seg 训练与调参、阈值与指标分析",
                "src/train.py、src/val.py、docs/train.md、docs/val.md",
            ],
            [
                "数据与标注",
                "队员B",
                "Label Studio 多边形标注、数据划分与增量整理、样本质量抽检",
                "src/organize_yolo_data.py、configs/data.yaml、data/yolo",
            ],
            [
                "推理与接口",
                "队员C",
                "批量推理、Mask 转外接框、八列 CSV、与附件1对齐",
                "src/predict.py、docs/接口规范.md、docs/predict.md",
            ],
            [
                "可视化应用",
                "队员D",
                "桌面壳、拖放导入、列表多选/全选/删除、选中导出与清空",
                "app/desktop.py、app/static/*",
            ],
            [
                "测试与文档",
                "队员E",
                "功能测试、导出用例核对、方案报告与演示材料",
                "本报告、app/README.md、提交前核对清单",
            ],
        ],
        [2.8, 2.4, 5.4, 5.2],
    )
    add_caption(doc, "表9 五人小组分工（姓名请按实际报名信息替换队员A–E）")

    add_heading(doc, "6.3 协作机制", 2)
    add_body(
        doc,
        "算法侧产出 best.pt 与验证指标后，由推理接口负责人固化 CONF 与 CSV 规则；应用侧只消费同一套字段，避免界面类别与提交文件不一致。重大规则（类别名、占位行、坐标定义）以 docs/接口规范.md 为单一事实来源。代码协作分支示例：no-image-enhancement 用于桌面交互增强（拖放、多选导出、选中清空等）。",
    )

    add_heading(doc, "七、总结与后续工作", 1)
    add_body(
        doc,
        "本方案面向绍兴市第一届大学生人工智能应用竞赛赛题二，完成了从工业背景理解、实例分割建模、数据整理、训练验证、规范导出到桌面可视化的完整闭环。验证集参考指标 Box mAP50 约 0.523，说明矩形定位已具备可用基础；Mask 高 IoU 指标偏低、训练集规模偏小，是下一阶段的主要改进点。",
    )
    add_body(
        doc,
        "后续工作建议：持续扩充多边形标注并重训；在 val 上系统扫描置信度—F1 曲线以确定提交阈值；将桌面应用由 Mock 检测切换为调用同一套 YOLO 权重，实现“所见即所交”；补充误检/漏检案例集，针对涂覆纹理与细裂纹做增强。所有改进仍须遵守“不碰测试集标签、不人工改写结果”的竞赛纪律。",
    )

    add_heading(doc, "参考文献与资料", 1)
    refs = [
        "[1] 绍兴市第一届大学生人工智能应用竞赛赛题：智能制造与工业智能 AI 应用赛道（赛题说明书）。",
        "[2] 附件1：检测结果输出表格规范及字段说明。",
        "[3] 仓库文档：docs/接口规范.md、docs/train.md、docs/val.md、docs/predict.md、docs/organize_yolo_data.md。",
        "[4] Ultralytics YOLO 官方文档：Train / Val / Predict / Segment 任务说明。",
        "[5] 仓库实现：src/train.py、src/val.py、src/predict.py、src/organize_yolo_data.py、app/desktop.py、configs/data.yaml。",
    ]
    for r in refs:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.74)
        p.paragraph_format.first_line_indent = Cm(-0.74)
        p.paragraph_format.line_spacing = 1.5
        run = p.add_run(r)
        set_run_font(run, "宋体", 10.5, False, INK, "宋体")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
